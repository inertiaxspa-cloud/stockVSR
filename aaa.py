import streamlit as st
import pandas as pd
import altair as alt
import io
from datetime import date
from streamlit_gsheets import GSheetsConnection
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Gestor de Sobre-Stock", layout="wide")

st.title("🍷 Monitor de Sobre-Stock e Inventario Inmovilizado")
st.write("Sube los reportes semanales completos en Excel para detectar oportunidades de movimiento de inventario.")

# --- INICIAR CONEXIÓN A GOOGLE SHEETS ---
try:
    conn = st.connection("gsheets", type=GSheetsConnection)
except Exception:
    conn = None


# ─────────────────────────────────────────────────────────────────────────────
# FUNCIÓN: Leer archivo Excel/CSV con limpieza defensiva
# ─────────────────────────────────────────────────────────────────────────────
def leer_archivo(archivo):
    nombre = archivo.name.lower()
    if nombre.endswith('.xlsx') or nombre.endswith('.xls'):
        try:
            xls = pd.ExcelFile(archivo, engine='openpyxl')
        except Exception as e:
            raise ValueError(f"No se pudo abrir el archivo Excel '{archivo.name}': {e}")

        hojas = xls.sheet_names
        hoja_objetivo = next((h for h in hojas if 'planilla' in h.strip().lower()), None)
        if hoja_objetivo is None:
            raise ValueError(
                f"No se encontró la pestaña 'PLANILLA' en '{archivo.name}'.\n"
                f"Hojas disponibles: {', '.join(hojas)}"
            )
        df = pd.read_excel(archivo, sheet_name=hoja_objetivo, engine='openpyxl')
    else:
        try:
            df = pd.read_csv(archivo, encoding='latin-1', sep=';')
        except Exception:
            archivo.seek(0)
            try:
                df = pd.read_csv(archivo, encoding='latin-1', sep=',')
            except Exception as e:
                raise ValueError(f"No se pudo leer el CSV '{archivo.name}': {e}")

    # ── Limpiar espacios en nombres de columnas (ej: 'Estatus ' → 'Estatus')
    df.columns = df.columns.str.strip()

    # ── Eliminar filas sin código de material (filas vacías al final del Excel)
    filas_totales = len(df)
    df = df.dropna(subset=['Material']).reset_index(drop=True)
    filas_descartadas = filas_totales - len(df)

    return df, filas_descartadas


# ─────────────────────────────────────────────────────────────────────────────
# FUNCIÓN: Formato de moneda
# ─────────────────────────────────────────────────────────────────────────────
def formato_moneda(valor):
    return f"${valor:,.0f}".replace(",", ".")


# ─────────────────────────────────────────────────────────────────────────────
# FUNCIÓN: Aplicar formato condicional al Excel de descarga
# ─────────────────────────────────────────────────────────────────────────────
def dar_formato_excel(writer, df, sheet_name, col_variacion=None, col_totales=None):
    """Escribe un DataFrame a una hoja Excel con formato: header, anchos, colores."""
    df.to_excel(writer, index=False, sheet_name=sheet_name)
    ws = writer.sheets[sheet_name]

    # Estilos
    fill_header = PatternFill("solid", fgColor="4A235A")   # Morado viña
    fill_pos    = PatternFill("solid", fgColor="C6EFCE")   # Verde claro
    fill_neg    = PatternFill("solid", fgColor="FFC7CE")   # Rojo claro
    fill_total  = PatternFill("solid", fgColor="D9D9D9")   # Gris
    font_header = Font(bold=True, color="FFFFFF")
    font_total  = Font(bold=True)
    border_thin = Border(
        bottom=Side(border_style="thin"),
        top=Side(border_style="thin"),
    )
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Encabezados
    for cell in ws[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center

    # Ajustar anchos de columna
    for col_idx, col in enumerate(ws.columns, 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 45)

    # Colores condicionales en columna de variación
    if col_variacion and col_variacion in df.columns:
        col_idx = df.columns.get_loc(col_variacion) + 1
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            try:
                v = float(cell.value or 0)
                cell.fill = fill_pos if v > 0 else (fill_neg if v < 0 else cell.fill)
            except (ValueError, TypeError):
                pass

    # Fila de totales
    if col_totales:
        total_row = ws.max_row + 1
        for col_name in col_totales:
            if col_name in df.columns:
                col_idx = df.columns.get_loc(col_name) + 1
                total = df[col_name].sum()
                cell = ws.cell(row=total_row, column=col_idx, value=total)
                cell.fill = fill_total
                cell.font = font_total
                cell.border = border_thin
        # Etiqueta "TOTAL"
        ws.cell(row=total_row, column=1, value="TOTAL").font = font_total


# ─────────────────────────────────────────────────────────────────────────────
# CARGA DE ARCHIVOS
# ─────────────────────────────────────────────────────────────────────────────
col1, col2 = st.columns(2)
with col1:
    archivo_anterior = st.file_uploader("📂 Sube el Excel de la semana ANTERIOR", type=['csv', 'xlsx', 'xls'])
with col2:
    archivo_actual = st.file_uploader("📂 Sube el Excel de la semana ACTUAL", type=['csv', 'xlsx', 'xls'])

# ─────────────────────────────────────────────────────────────────────────────
# PROCESAMIENTO PRINCIPAL
# ─────────────────────────────────────────────────────────────────────────────
if archivo_anterior and archivo_actual:

    # ── PASO 1: Leer archivos ────────────────────────────────────────────────
    try:
        df_ant, desc_ant = leer_archivo(archivo_anterior)
        df_act, desc_act = leer_archivo(archivo_actual)
    except ValueError as e:
        st.error(f"❌ Error al leer los archivos:\n\n{e}")
        st.stop()
    except Exception as e:
        st.error(f"❌ Error inesperado al leer archivos: {e}")
        st.stop()

    # ── Indicador de calidad de datos ────────────────────────────────────────
    with st.expander("🔍 Calidad de los datos cargados", expanded=False):
        c1, c2 = st.columns(2)
        with c1:
            st.markdown(f"**Semana anterior:** `{archivo_anterior.name}`")
            st.write(f"- Filas válidas: **{len(df_ant):,}**")
            if desc_ant > 0:
                st.warning(f"⚠️ {desc_ant} filas vacías descartadas")
            almacenes_invalidos_ant = df_ant['Almacén'].astype(str).str.strip().str.upper().isin(
                ['FALSO', 'FALSE', 'NAN', '']) if 'Almacén' in df_ant.columns else pd.Series([False])
            if almacenes_invalidos_ant.sum() > 0:
                st.warning(f"⚠️ {almacenes_invalidos_ant.sum()} registros con Almacén inválido (False/FALSO)")
        with c2:
            st.markdown(f"**Semana actual:** `{archivo_actual.name}`")
            st.write(f"- Filas válidas: **{len(df_act):,}**")
            if desc_act > 0:
                st.warning(f"⚠️ {desc_act} filas vacías descartadas")
            almacenes_invalidos_act = df_act['Almacén'].astype(str).str.strip().str.upper().isin(
                ['FALSO', 'FALSE', 'NAN', '']) if 'Almacén' in df_act.columns else pd.Series([False])
            if almacenes_invalidos_act.sum() > 0:
                st.warning(f"⚠️ {almacenes_invalidos_act.sum()} registros con Almacén inválido (False/FALSO)")

    # ── PASO 2: Validar columnas requeridas ───────────────────────────────────
    COLS_REQUERIDAS = ['Material', 'LOTE', 'Texto breve de material',
                       'Libre utilización', 'Valor libre util.', 'Almacén', 'Estatus']
    for label, df_check in [("semana anterior", df_ant), ("semana actual", df_act)]:
        faltantes = [c for c in COLS_REQUERIDAS if c not in df_check.columns]
        if faltantes:
            st.error(
                f"❌ Al archivo de **{label}** le faltan las columnas: `{'`, `'.join(faltantes)}`\n\n"
                f"Columnas encontradas: `{'`, `'.join(df_check.columns.tolist())}`"
            )
            st.stop()

    # ── PASO 3: KPIs de No Vigente (solo sobre df_act limpio) ─────────────────
    try:
        df_act['Almacén'] = df_act['Almacén'].fillna('').astype(str).str.strip()
        df_act['Estatus'] = df_act['Estatus'].fillna('').astype(str).str.strip()
        df_act['Libre utilización'] = pd.to_numeric(df_act['Libre utilización'], errors='coerce').fillna(0)
        df_act['Valor libre util.'] = pd.to_numeric(df_act['Valor libre util.'], errors='coerce').fillna(0)

        valores_excluir = ['FALSO', 'FALSE', '#N/A', 'NAN', '']
        m_estatus = df_act['Estatus'].str.upper() == 'NO VIGENTE'
        m_almacen = ~df_act['Almacén'].str.upper().isin(valores_excluir)
        df_kpi_no_vigente = df_act[m_estatus & m_almacen]

        total_no_vigente = int(df_kpi_no_vigente['Libre utilización'].sum())
        valor_no_vigente = df_kpi_no_vigente['Valor libre util.'].sum()
        valor_total_act  = df_act['Valor libre util.'].sum()
    except Exception as e:
        st.error(f"❌ Error al calcular KPIs de No Vigente: {e}")
        st.stop()

    # ── PASO 4: Preparar columnas clave y normalizar ANTES del merge ──────────
    try:
        columnas_clave = ['Material', 'LOTE', 'Texto breve de material',
                          'Libre utilización', 'Valor libre util.', 'Almacén', 'Estatus']

        # Asegurar que columnas faltantes existan (con tipo correcto)
        COLS_NUMERICAS = {'Libre utilización', 'Valor libre util.'}
        for col in columnas_clave:
            if col not in df_ant.columns:
                df_ant[col] = 0 if col in COLS_NUMERICAS else ''
            if col not in df_act.columns:
                df_act[col] = 0 if col in COLS_NUMERICAS else ''

        # Normalizar claves y tipos ANTES del merge
        for df_tmp in [df_ant, df_act]:
            df_tmp['Material']                = df_tmp['Material'].astype(str).str.strip()
            df_tmp['LOTE']                    = df_tmp['LOTE'].astype(str).str.strip().str.upper().replace('NAN', '')
            df_tmp['Texto breve de material'] = df_tmp['Texto breve de material'].fillna('').astype(str).str.strip()
            df_tmp['Libre utilización']       = pd.to_numeric(df_tmp['Libre utilización'], errors='coerce').fillna(0)
            df_tmp['Valor libre util.']       = pd.to_numeric(df_tmp['Valor libre util.'], errors='coerce').fillna(0)
            df_tmp['Almacén']                 = df_tmp['Almacén'].fillna('').astype(str).str.strip()
            df_tmp['Estatus']                 = df_tmp['Estatus'].fillna('').astype(str).str.strip()
    except Exception as e:
        st.error(f"❌ Error al normalizar los datos antes del cruce: {e}")
        st.stop()

    # ── PASO 5: Merge outer y cálculo de variaciones ─────────────────────────
    try:
        df_cruce = pd.merge(
            df_ant[columnas_clave], df_act[columnas_clave],
            on=['Material', 'LOTE', 'Texto breve de material'],
            suffixes=('_Ant', '_Act'),
            how='outer'
        )

        # fillna inteligente: texto → '' | número → 0
        for _col in df_cruce.columns:
            if df_cruce[_col].dtype == 'object':
                df_cruce[_col] = df_cruce[_col].fillna('')
            else:
                df_cruce[_col] = df_cruce[_col].fillna(0)

        # Forzar tipos numéricos en columnas de cálculo (por si quedaron mixed)
        for c in ['Libre utilización_Ant', 'Libre utilización_Act',
                  'Valor libre util._Ant', 'Valor libre util._Act']:
            df_cruce[c] = pd.to_numeric(df_cruce[c], errors='coerce').fillna(0)

        df_cruce['Variacion_Unidades'] = df_cruce['Libre utilización_Act'] - df_cruce['Libre utilización_Ant']
        df_cruce['Variacion_Valor']    = df_cruce['Valor libre util._Act'] - df_cruce['Valor libre util._Ant']

    except Exception as e:
        st.error(f"❌ Error al cruzar los datos entre semanas: {e}")
        st.stop()

    # ── PASO 6: Clasificación y métricas derivadas ────────────────────────────
    try:
        def determinar_estado(row):
            if row['Libre utilización_Ant'] == 0 and row['Libre utilización_Act'] > 0:
                return "🆕 Material Nuevo"
            if row['Libre utilización_Act'] == 0 and row['Libre utilización_Ant'] > 0:
                return "🚫 Desaparecido"
            return "Ya Estaba"

        def calcular_porcentaje(row):
            ant = float(row['Libre utilización_Ant']) if row['Libre utilización_Ant'] != '' else 0.0
            var = float(row['Variacion_Unidades'])    if row['Variacion_Unidades']    != '' else 0.0
            if var <= 0:
                return "0%"
            if ant == 0:
                return "100% (Nuevo)"
            return f"+{(var / ant) * 100:.1f}%"

        df_cruce['Estado Material'] = df_cruce.apply(determinar_estado, axis=1)
        df_cruce['% Aumento']       = df_cruce.apply(calcular_porcentaje, axis=1)

        # Lote limpio para gráficos (sin 'nan' literal)
        df_cruce['LOTE_display']    = df_cruce['LOTE'].replace({'': '—', 'NAN': '—'})
        df_cruce['Nombre_Grafico']  = (
            df_cruce['Texto breve de material'].astype(str)
            + " (Lote: " + df_cruce['LOTE_display'] + ")"
        )

        # Segmentaciones
        sobre_stock       = df_cruce[(df_cruce['Variacion_Unidades'] > 0) |
                                     ((df_cruce['Variacion_Unidades'] == 0) & (df_cruce['Libre utilización_Act'] > 500))].copy()
        solo_aumentos     = df_cruce[df_cruce['Variacion_Unidades'] > 0].copy()
        solo_bajadas      = df_cruce[df_cruce['Variacion_Unidades'] < 0].copy()
        materiales_nuevos = df_cruce[df_cruce['Estado Material'] == '🆕 Material Nuevo'].copy()
        desaparecidos     = df_cruce[df_cruce['Estado Material'] == '🚫 Desaparecido'].copy()

    except Exception as e:
        st.error(f"❌ Error al clasificar los datos: {e}")
        st.stop()

    # ─────────────────────────────────────────────────────────────────────────
    # VISUALIZACIÓN
    # ─────────────────────────────────────────────────────────────────────────
    if not df_cruce.empty:
        st.divider()
        tab1, tab2, tab3 = st.tabs(["📊 Dashboard Visual", "🔍 Reportes y Descargas", "☁️ Trazabilidad Histórica"])

        # ══════════════════════════════════════════════════════════════════════
        with tab1:
            st.header("Dashboard Ejecutivo de Inventario")

            unidades_subieron  = int(solo_aumentos['Variacion_Unidades'].sum())
            unidades_bajaron   = int(solo_bajadas['Variacion_Unidades'].sum())
            valor_ingresado    = solo_aumentos['Variacion_Valor'].sum()
            pct_inmovilizado   = (valor_no_vigente / valor_total_act * 100) if valor_total_act > 0 else 0

            # ── Fila 1 de métricas ────────────────────────────────────────────
            m1, m2, m3, m4, m5 = st.columns(5)
            m1.metric("📈 Mat. que SUBIERON",     len(solo_aumentos),
                      help="Materiales con más stock que la semana anterior")
            m2.metric("📉 Mat. que BAJARON",      len(solo_bajadas),
                      help="Materiales con menos stock que la semana anterior")
            m3.metric("🆕 Materiales Nuevos",     len(materiales_nuevos),
                      help="No existían la semana anterior")
            m4.metric("🚫 Desaparecidos",         len(desaparecidos),
                      help="Tenían stock y ahora están en 0")
            m5.metric("📦 Unidades Ingresadas",   f"{unidades_subieron:,}".replace(",", "."))

            st.write("---")

            # ── Fila 2 de métricas ────────────────────────────────────────────
            k1, k2, k3, k4 = st.columns(4)
            k1.metric("📉 Unidades que Bajaron",      f"{abs(unidades_bajaron):,}".replace(",", "."),
                      delta=f"{unidades_bajaron:,}".replace(",", "."), delta_color="inverse")
            k2.metric("💰 Capital Ingresado",          formato_moneda(valor_ingresado))
            k3.metric("⚠️ Unid. 'No Vigentes'",        f"{total_no_vigente:,}".replace(",", "."))
            k4.metric("🏦 Capital 'No Vigente'",        formato_moneda(valor_no_vigente),
                      help=f"{pct_inmovilizado:.1f}% del capital total del inventario actual")

            # Barra de % capital inmovilizado
            st.caption(f"📊 El capital No Vigente representa el **{pct_inmovilizado:.1f}%** del inventario total actual")
            st.progress(min(pct_inmovilizado / 100, 1.0))

            st.write("---")

            # ── Gráficos: Top Subidas vs Top Bajadas ──────────────────────────
            grafico_izq, grafico_der = st.columns(2)

            with grafico_izq:
                st.subheader("📈 Top 10 Mayores Subidas")
                top_subidas = solo_aumentos.sort_values('Variacion_Unidades', ascending=False).head(10).copy()
                if not top_subidas.empty:
                    top_subidas['Etiqueta'] = top_subidas['Variacion_Unidades'].apply(
                        lambda x: f"+{int(x):,}".replace(',', '.'))
                    bars = alt.Chart(top_subidas).mark_bar(color='#E15A97').encode(
                        x=alt.X('Variacion_Unidades:Q', title='Unidades ingresadas'),
                        y=alt.Y('Nombre_Grafico:N', sort='-x', title='', axis=alt.Axis(labelLimit=300))
                    )
                    text = bars.mark_text(align='left', dx=4, fontWeight='bold').encode(
                        text=alt.Text('Etiqueta:N'))
                    st.altair_chart((bars + text).properties(height=350), use_container_width=True)
                else:
                    st.info("No hubo subidas esta semana.")

            with grafico_der:
                st.subheader("📉 Top 10 Mayores Bajadas")
                top_bajadas = solo_bajadas.sort_values('Variacion_Unidades', ascending=True).head(10).copy()
                if not top_bajadas.empty:
                    top_bajadas['Etiqueta'] = top_bajadas['Variacion_Unidades'].apply(
                        lambda x: f"{int(x):,}".replace(',', '.'))
                    bars_b = alt.Chart(top_bajadas).mark_bar(color='#4A90E2').encode(
                        x=alt.X('Variacion_Unidades:Q', title='Unidades reducidas'),
                        y=alt.Y('Nombre_Grafico:N', sort='x', title='', axis=alt.Axis(labelLimit=300))
                    )
                    text_b = bars_b.mark_text(align='right', dx=-4, fontWeight='bold').encode(
                        text=alt.Text('Etiqueta:N'))
                    st.altair_chart((bars_b + text_b).properties(height=350), use_container_width=True)
                else:
                    st.success("✅ No hubo bajadas de stock esta semana.")

            # ── Materiales nuevos y desaparecidos ─────────────────────────────
            st.write("---")
            col_nuevos, col_desap = st.columns(2)
            with col_nuevos:
                st.subheader("🆕 Materiales Nuevos")
                if not materiales_nuevos.empty:
                    st.dataframe(
                        materiales_nuevos[['Material', 'Texto breve de material', 'LOTE',
                                           'Almacén_Act', 'Libre utilización_Act', 'Valor libre util._Act']]
                        .rename(columns={'Libre utilización_Act': 'Stock Actual',
                                         'Valor libre util._Act': 'Valor ($)',
                                         'Almacén_Act': 'Almacén'}),
                        use_container_width=True, hide_index=True
                    )
                else:
                    st.info("No hay materiales nuevos esta semana.")

            with col_desap:
                st.subheader("🚫 Materiales Desaparecidos")
                if not desaparecidos.empty:
                    st.dataframe(
                        desaparecidos[['Material', 'Texto breve de material', 'LOTE',
                                       'Almacén_Ant', 'Libre utilización_Ant', 'Valor libre util._Ant']]
                        .rename(columns={'Libre utilización_Ant': 'Stock Anterior',
                                         'Valor libre util._Ant': 'Valor Anterior ($)',
                                         'Almacén_Ant': 'Almacén'}),
                        use_container_width=True, hide_index=True
                    )
                else:
                    st.success("✅ Ningún material desapareció esta semana.")

        # ══════════════════════════════════════════════════════════════════════
        with tab2:
            # ── Auditoría KPI ─────────────────────────────────────────────────
            st.subheader("🕵️‍♂️ Auditoría de KPI: ¿Qué se sumó en 'No Vigente'?")
            with st.expander("Ver lista de materiales contabilizados"):
                st.write(f"**Total filas:** {len(df_kpi_no_vigente):,}")
                st.write("**Filtros:** Estatus = 'NO VIGENTE' | Almacén válido (no FALSO, no vacío)")
                st.dataframe(
                    df_kpi_no_vigente[['Material', 'Almacén', 'Estatus', 'Libre utilización', 'Valor libre util.']]
                    .rename(columns={'Libre utilización': 'Unidades', 'Valor libre util.': 'Valor ($)'}),
                    use_container_width=True, hide_index=True
                )

            st.write("---")

            # ── Reporte de Aumentos ───────────────────────────────────────────
            st.subheader("📈 Reporte de Aumentos de Inventario")
            cols_aumentos_vista = ['Material', 'Estado Material', 'Almacén_Act', 'LOTE',
                                   'Texto breve de material', 'Libre utilización_Ant',
                                   'Libre utilización_Act', 'Variacion_Unidades', '% Aumento']
            st.dataframe(
                solo_aumentos.sort_values('Variacion_Unidades', ascending=False)[cols_aumentos_vista]
                .rename(columns={'Libre utilización_Ant': 'Stock Anterior',
                                 'Libre utilización_Act': 'Stock Actual',
                                 'Variacion_Unidades':    'Diferencia (+)',
                                 'Almacén_Act':           'Almacén'}),
                use_container_width=True, hide_index=True,
                column_config={
                    "Stock Anterior": st.column_config.NumberColumn(format="%d"),
                    "Stock Actual":   st.column_config.NumberColumn(format="%d"),
                    "Diferencia (+)": st.column_config.NumberColumn(format="%d"),
                }
            )

            st.write("---")

            # ── Plan de Acción ────────────────────────────────────────────────
            st.subheader("📋 Detalle General y Plan de Acción")
            if not sobre_stock.empty:
                sobre_stock_sorted = sobre_stock.sort_values('Libre utilización_Act', ascending=False).copy()

                def generar_recomendacion(row):
                    if row['Variacion_Unidades'] > 500:
                        return "🔴 Alerta: Fuerte ingreso. Confirmar justificación."
                    elif row['Variacion_Unidades'] > 0:
                        return "🟡 Aumento de stock. Vigilar rotación."
                    elif row['Libre utilización_Act'] > 5000:
                        return "🔵 Inmovilizado Alto: Evaluar Venta Ecommerce."
                    elif row['Libre utilización_Act'] > 1000:
                        return "🟢 Inmovilizado Medio: Sugerir Solicitudes Turismo."
                    else:
                        return "⚪ Inmovilizado Bajo: Armar packs promocionales."

                sobre_stock_sorted['Recomendación'] = sobre_stock_sorted.apply(generar_recomendacion, axis=1)

                cols_plan = ['Material', 'Estado Material', 'Almacén_Act', 'LOTE',
                             'Texto breve de material', 'Libre utilización_Act',
                             'Variacion_Unidades', 'Valor libre util._Act', 'Recomendación']
                st.dataframe(
                    sobre_stock_sorted[cols_plan]
                    .rename(columns={'Libre utilización_Act': 'Stock Actual',
                                     'Variacion_Unidades':    'Variación (Unid.)',
                                     'Valor libre util._Act': 'Valor Actual ($)',
                                     'Almacén_Act':           'Almacén Actual'}),
                    use_container_width=True, hide_index=True,
                    column_config={
                        "Valor Actual ($)": st.column_config.NumberColumn(format="$ %d"),
                        "Stock Actual":     st.column_config.NumberColumn(format="%d"),
                        "Variación (Unid.)":st.column_config.NumberColumn(format="%d"),
                    }
                )
            else:
                st.success("No hay alertas de inventario esta semana.")

            st.divider()

            # ── GENERACIÓN DEL EXCEL DE DESCARGA ──────────────────────────────
            st.subheader("📥 Descargar Reporte Completo")

            # Columnas en español para el Excel
            COLS_EXPORT = {
                'Material':                  'Material',
                'LOTE':                      'Lote',
                'Texto breve de material':   'Descripción',
                'Almacén_Ant':               'Almacén Anterior',
                'Almacén_Act':               'Almacén Actual',
                'Estatus_Ant':               'Estatus Anterior',
                'Estatus_Act':               'Estatus Actual',
                'Libre utilización_Ant':     'Stock Semana Anterior',
                'Libre utilización_Act':     'Stock Semana Actual',
                'Variacion_Unidades':        'Diferencia Unidades',
                'Variacion_Valor':           'Diferencia Valor ($)',
                'Valor libre util._Ant':     'Valor Anterior ($)',
                'Valor libre util._Act':     'Valor Actual ($)',
                'Estado Material':           'Estado',
                '% Aumento':                 '% Cambio',
            }

            def prep_export(df_in, cols_subset=None):
                cols = cols_subset or list(COLS_EXPORT.keys())
                available = [c for c in cols if c in df_in.columns]
                return df_in[available].rename(columns=COLS_EXPORT).copy()

            output = io.BytesIO()
            with pd.ExcelWriter(output, engine='openpyxl') as writer:

                # Hoja 1: Subidas
                df_exp_subidas = prep_export(
                    solo_aumentos.sort_values('Variacion_Unidades', ascending=False),
                    ['Material', 'Texto breve de material', 'LOTE', 'Almacén_Act',
                     'Libre utilización_Ant', 'Libre utilización_Act', 'Variacion_Unidades',
                     '% Aumento', 'Valor libre util._Act', 'Estado Material']
                )
                dar_formato_excel(writer, df_exp_subidas, '📈 Subidas de Stock',
                                  col_variacion='Diferencia Unidades',
                                  col_totales=['Stock Semana Anterior', 'Stock Semana Actual',
                                               'Diferencia Unidades', 'Valor Actual ($)'])

                # Hoja 2: Bajadas
                df_exp_bajadas = prep_export(
                    solo_bajadas.sort_values('Variacion_Unidades', ascending=True),
                    ['Material', 'Texto breve de material', 'LOTE', 'Almacén_Act',
                     'Libre utilización_Ant', 'Libre utilización_Act', 'Variacion_Unidades',
                     'Valor libre util._Ant', 'Estado Material']
                )
                dar_formato_excel(writer, df_exp_bajadas, '📉 Bajadas de Stock',
                                  col_variacion='Diferencia Unidades',
                                  col_totales=['Stock Semana Anterior', 'Stock Semana Actual',
                                               'Diferencia Unidades'])

                # Hoja 3: Materiales nuevos
                df_exp_nuevos = prep_export(
                    materiales_nuevos,
                    ['Material', 'Texto breve de material', 'LOTE', 'Almacén_Act',
                     'Libre utilización_Act', 'Valor libre util._Act']
                )
                dar_formato_excel(writer, df_exp_nuevos, '🆕 Materiales Nuevos',
                                  col_totales=['Stock Semana Actual', 'Valor Actual ($)'])

                # Hoja 4: Desaparecidos
                df_exp_desap = prep_export(
                    desaparecidos,
                    ['Material', 'Texto breve de material', 'LOTE', 'Almacén_Ant',
                     'Libre utilización_Ant', 'Valor libre util._Ant']
                )
                dar_formato_excel(writer, df_exp_desap, '🚫 Desaparecidos',
                                  col_totales=['Stock Semana Anterior', 'Valor Anterior ($)'])

                # Hoja 5: Cruce Completo
                df_exp_total = prep_export(df_cruce.sort_values('Material'))
                dar_formato_excel(writer, df_exp_total, '📋 Cruce Completo',
                                  col_variacion='Diferencia Unidades',
                                  col_totales=['Stock Semana Anterior', 'Stock Semana Actual',
                                               'Diferencia Unidades'])

                # Hoja 6: No Vigentes
                df_exp_nv = df_kpi_no_vigente[
                    [c for c in ['Material', 'Texto breve de material', 'LOTE',
                                 'Almacén', 'Estatus', 'Libre utilización', 'Valor libre util.']
                     if c in df_kpi_no_vigente.columns]
                ].rename(columns={'Libre utilización': 'Unidades',
                                  'Valor libre util.': 'Valor ($)'}).copy()
                dar_formato_excel(writer, df_exp_nv, '⚠️ Alertas No Vigente',
                                  col_totales=['Unidades', 'Valor ($)'])

            st.download_button(
                label="📥 Descargar Reporte Completo (6 hojas)",
                data=output.getvalue(),
                file_name="Reporte_Inventario.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary"
            )

        # ══════════════════════════════════════════════════════════════════════
        with tab3:
            st.header("Base de Datos Histórica (Google Sheets)")
            if conn is not None:
                with st.form("form_guardar_bd"):
                    col_fecha, col_btn = st.columns([1, 2])
                    with col_fecha:
                        fecha_registro = st.date_input("Fecha de esta foto de inventario:", date.today())
                    with col_btn:
                        st.write(""); st.write("")
                        guardar = st.form_submit_button("💾 Enviar 'Semana Actual' a Google Sheets")

                    if guardar:
                        with st.spinner("Conectando con Google Sheets..."):
                            try:
                                df_hist = conn.read(worksheet="Historial", usecols=list(range(6)), ttl=0).dropna(how="all")
                            except Exception:
                                df_hist = pd.DataFrame(columns=['Fecha_Registro', 'Material', 'LOTE',
                                                                 'Texto_breve', 'Libre_utilizacion', 'Valor'])

                            cols_bd = ['Material', 'LOTE', 'Texto breve de material',
                                       'Libre utilización', 'Valor libre util.']
                            faltantes_bd = [c for c in cols_bd if c not in df_act.columns]
                            if faltantes_bd:
                                st.error(f"Faltan columnas para guardar: {faltantes_bd}")
                            else:
                                df_para_bd = df_act[cols_bd].copy()
                                df_para_bd.rename(columns={
                                    'Texto breve de material': 'Texto_breve',
                                    'Libre utilización':       'Libre_utilizacion',
                                    'Valor libre util.':       'Valor'
                                }, inplace=True)
                                df_para_bd.insert(0, 'Fecha_Registro', str(fecha_registro))

                                if not df_hist.empty:
                                    df_hist['Fecha_Registro'] = df_hist['Fecha_Registro'].astype(str)
                                    df_hist = df_hist[df_hist['Fecha_Registro'] != str(fecha_registro)]

                                df_updated = pd.concat([df_hist, df_para_bd], ignore_index=True)
                                conn.update(worksheet="Historial", data=df_updated)
                                st.success(f"✅ Inventario del {fecha_registro} guardado correctamente en Google Sheets.")

                st.divider()
                st.subheader("📈 Análisis de Tendencias Históricas")
                if st.button("🔄 Cargar Gráficos Históricos"):
                    with st.spinner("Descargando historial desde Google..."):
                        try:
                            df_hist_cloud = conn.read(worksheet="Historial", usecols=list(range(6)), ttl=0).dropna(how="all")
                            if not df_hist_cloud.empty:
                                df_hist_cloud['Fecha_Registro'] = pd.to_datetime(df_hist_cloud['Fecha_Registro'], errors='coerce')
                                material_sel = st.selectbox("Selecciona un material:",
                                                            sorted(df_hist_cloud['Texto_breve'].dropna().unique()))
                                datos_grafico = df_hist_cloud[df_hist_cloud['Texto_breve'] == material_sel].copy()
                                if not datos_grafico.empty:
                                    datos_grafico['Libre_utilizacion'] = pd.to_numeric(datos_grafico['Libre_utilizacion'], errors='coerce')
                                    linea = alt.Chart(datos_grafico).mark_line(point=True, color='#FF5722', strokeWidth=3).encode(
                                        x=alt.X('Fecha_Registro:T', title='Fecha'),
                                        y=alt.Y('Libre_utilizacion:Q', title='Stock Total (Unidades)'),
                                        color=alt.Color('LOTE:N', legend=alt.Legend(title="Lotes"))
                                    )
                                    st.altair_chart(linea.properties(height=400), use_container_width=True)
                            else:
                                st.info("Aún no has guardado ningún dato histórico.")
                        except Exception as e:
                            st.error(f"No se pudo cargar el historial: {e}")
            else:
                st.warning("⚠️ La conexión a Google Sheets no está configurada. Los datos solo estarán disponibles en el reporte Excel.")
